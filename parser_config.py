"""
Parser and config
=================
Reads the Excel rulebook, attaches the hand-written rule specs and the
user-defined intervals, validates everything and writes one versioned,
machine-readable file: config/parameters_config.json

It does NOT touch entity data, normalize values, or calculate scores.

Run from anywhere:
    python parser_config.py                    # build the config
    python parser_config.py --strict           # refuse to build if any interval is missing
    python parser_config.py --init-intervals   # create an empty intervals.json template
"""
from __future__ import annotations

import argparse
import hashlib
import json
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Optional

from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parent
CONFIG_DIR = BASE_DIR / "config"
DEFAULT_EXCEL = CONFIG_DIR / "Due_Dilligence_Weighed_Parameters.xlsx"
DEFAULT_RULE_SPECS = CONFIG_DIR / "rule_specs.json"
DEFAULT_INTERVALS = CONFIG_DIR / "intervals.json"
DEFAULT_OUTPUT = CONFIG_DIR / "parameters_config.json"
DEFAULT_VERSIONS_DIR = CONFIG_DIR / "versions"

# The workbook says "Weigh"; internally we call it "weight".
HEADER_ALIASES = {
    "parameter": "parameter",
    "field": "field",
    "context": "context",
    "rule": "rule",
    "weigh": "weight",
    "weight": "weight",
}
REQUIRED_COLUMNS = ("parameter", "field", "context", "rule", "weight")

RULE_TYPES = ("banded_then_decay", "base_plus_terms", "categorical", "boolean_map")
DERIVED_OPS = ("months_between",)


class ConfigError(Exception):
    """Raised for any problem with the source document or configuration files."""


# --------------------------------------------------------------------------
# Small helpers
# --------------------------------------------------------------------------
def _is_number(value: Any) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def _clean(value: Any) -> Any:
    """Strip surrounding whitespace from text; empty text becomes None."""
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


def file_sha256(path: Path) -> str:
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def _read_json(path: Path, what: str) -> Any:
    path = Path(path)
    if not path.exists():
        raise ConfigError(f"{what} not found: {path}")
    try:
        with path.open("r", encoding="utf-8") as handle:
            return json.load(handle)
    except json.JSONDecodeError as exc:
        raise ConfigError(f"{what} is not valid JSON ({path}): {exc}") from exc


def _raise_if(errors: list[str], title: str) -> None:
    if errors:
        raise ConfigError(title + ":\n  - " + "\n  - ".join(errors))


def _version_of(parameters: list[dict]) -> str:
    """Content-based version: any change to a rule, weight or interval changes it."""
    canonical = json.dumps(parameters, sort_keys=True, ensure_ascii=False, separators=(",", ":"))
    return hashlib.sha256(canonical.encode("utf-8")).hexdigest()[:12]


# --------------------------------------------------------------------------
# 1. Parse the Excel document
# --------------------------------------------------------------------------
def _find_header(rows: list[tuple]) -> tuple[int, dict[str, int]]:
    for index, row in enumerate(rows[:10]):
        mapping: dict[str, int] = {}
        for column, value in enumerate(row):
            if isinstance(value, str):
                key = HEADER_ALIASES.get(value.strip().lower())
                if key and key not in mapping:
                    mapping[key] = column
        if all(name in mapping for name in REQUIRED_COLUMNS):
            return index, mapping
    raise ConfigError(
        "Could not find the header row (Parameter, Field, Context, Rule, Weigh/Weight) "
        "in the first 10 rows of the sheet."
    )


def _to_weight(value: Any, parameter: str, row_number: int) -> float | int:
    if isinstance(value, str):
        try:
            value = float(value.replace(",", "."))
        except ValueError:
            pass
    if not _is_number(value):
        raise ConfigError(f"Row {row_number} ('{parameter}'): weight is missing or not numeric ({value!r}).")
    if value < 0:
        raise ConfigError(f"Row {row_number} ('{parameter}'): weight cannot be negative ({value}).")
    return int(value) if float(value).is_integer() else float(value)


def parse_excel(path: Path | str = DEFAULT_EXCEL, sheet_name: Optional[str] = None) -> dict:
    """
    Extract the parameter table. Parameter, Field, Context and Rule are kept as
    written (only surrounding whitespace is trimmed), including typos.
    IDs (p01, p02, ...) follow row order.
    """
    path = Path(path)
    if not path.exists():
        raise ConfigError(f"Excel file not found: {path}")

    # NOTE: deliberately NOT read_only=True. In read-only mode openpyxl returns
    # max_row/max_column = None for this workbook, which crashed the first version.
    workbook = load_workbook(filename=path, data_only=True)
    try:
        if sheet_name is None:
            sheet = workbook.worksheets[0]
        elif sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
        else:
            raise ConfigError(f"Sheet '{sheet_name}' not found. Available: {', '.join(workbook.sheetnames)}")
        rows = list(sheet.iter_rows(values_only=True))
        sheet_title = sheet.title
    finally:
        workbook.close()

    header_index, columns = _find_header(rows)

    note = None
    for row in rows[:header_index]:
        for value in row:
            if _clean(value) is not None:
                note = str(_clean(value))
                break
        if note:
            break

    parameters: list[dict] = []
    for offset, row in enumerate(rows[header_index + 1:], start=header_index + 2):
        cells = {name: _clean(row[col]) if col < len(row) else None for name, col in columns.items()}
        if all(value is None for value in cells.values()):
            continue  # empty padding row
        if cells["parameter"] is None:
            raise ConfigError(f"Row {offset}: has content but no Parameter name.")
        parameters.append({
            "id": f"p{len(parameters) + 1:02d}",
            "parameter": str(cells["parameter"]),
            "field": "" if cells["field"] is None else str(cells["field"]),
            "context": "" if cells["context"] is None else str(cells["context"]),
            "rule": "" if cells["rule"] is None else str(cells["rule"]),
            "weight": _to_weight(cells["weight"], str(cells["parameter"]), offset),
        })

    if not parameters:
        raise ConfigError("No parameter rows were found below the header.")
    if sum(p["weight"] for p in parameters) <= 0:
        raise ConfigError("The total weight is 0; at least one parameter needs a positive weight.")
    for p in parameters:
        if not p["field"] or not p["rule"]:
            raise ConfigError(f"{p['id']} ('{p['parameter']}'): Field and Rule cannot be empty.")

    return {
        "source_file": path.name,
        "source_sha256": file_sha256(path),
        "sheet": sheet_title,
        "source_note": note,
        "parameters": parameters,
    }


# --------------------------------------------------------------------------
# 2. Rule specs (machine-executable version of each Rule sentence)
# --------------------------------------------------------------------------
def _validate_spec(pid: str, spec: dict, errors: list[str]) -> None:
    def err(message: str) -> None:
        errors.append(f"{pid}: {message}")

    def unit(value: Any, where: str) -> None:
        if not (_is_number(value) and 0 <= value <= 1):
            err(f"{where} must be a number between 0 and 1 (got {value!r})")

    spec_type = spec.get("type")
    if spec_type not in RULE_TYPES:
        err(f"unknown rule type {spec_type!r} (allowed: {', '.join(RULE_TYPES)})")
        return

    required = spec.get("required_facts")
    if not (isinstance(required, list) and required and all(isinstance(x, str) and x for x in required)):
        err("'required_facts' must be a non-empty list of fact names")
        return
    if len(set(required)) != len(required):
        err("'required_facts' contains duplicates")

    known = set(required)
    derived = spec.get("derived_inputs", {})
    if not isinstance(derived, dict):
        err("'derived_inputs' must be an object")
        derived = {}
    for name, definition in derived.items():
        if not isinstance(definition, dict) or definition.get("op") not in DERIVED_OPS:
            err(f"derived input '{name}' needs an 'op' from: {', '.join(DERIVED_OPS)}")
            continue
        if definition.get("from") not in required:
            err(f"derived input '{name}': 'from' must be listed in required_facts")
        if definition.get("to") != "checked_at" and definition.get("to") not in required:
            err(f"derived input '{name}': 'to' must be 'checked_at' or a required fact")
        known.add(name)

    def need(fact: Any, where: str) -> None:
        if fact not in known:
            err(f"{where} refers to '{fact}', which is not in required_facts or derived_inputs")

    if spec_type == "banded_then_decay":
        need(spec.get("on"), "'on'")
        bands = spec.get("bands")
        previous_end: float = 0
        if not (isinstance(bands, list) and bands):
            err("'bands' must be a non-empty list of [from, to, value]")
        else:
            for i, band in enumerate(bands, 1):
                if not (isinstance(band, list) and len(band) == 3 and all(_is_number(x) for x in band)):
                    err(f"band {i} must be [from, to, value]")
                    continue
                low, high, value = band
                if low != previous_end:
                    err(f"band {i} starts at {low} but should start at {previous_end} (bands must be contiguous from 0)")
                if high <= low:
                    err(f"band {i}: 'to' must be greater than 'from'")
                unit(value, f"band {i} value")
                previous_end = high
        after = spec.get("after")
        if not isinstance(after, dict):
            err("'after' is required (start_at, base, step_size, per_step, floor)")
        else:
            for key in ("start_at", "base", "step_size", "per_step", "floor"):
                if not _is_number(after.get(key)):
                    err(f"'after.{key}' must be a number")
            if _is_number(after.get("start_at")) and after["start_at"] != previous_end:
                err(f"'after.start_at' ({after['start_at']}) must equal the end of the last band ({previous_end})")
            if _is_number(after.get("step_size")) and after["step_size"] <= 0:
                err("'after.step_size' must be greater than 0")
            for key in ("base", "floor"):
                if _is_number(after.get(key)):
                    unit(after[key], f"'after.{key}'")

    elif spec_type == "base_plus_terms":
        unit(spec.get("base"), "'base'")
        if spec.get("clip", [0, 1]) != [0, 1]:
            err("'clip' must be [0, 1]")
        terms = spec.get("terms")
        if not (isinstance(terms, list) and terms):
            err("'terms' must be a non-empty list")
        else:
            for i, term in enumerate(terms, 1):
                if not isinstance(term, dict):
                    err(f"term {i} must be an object")
                    continue
                need(term.get("fact"), f"term {i}")
                kinds = [k for k in ("per_unit", "contains", "equals") if k in term]
                if len(kinds) != 1:
                    err(f"term {i} must have exactly one of: per_unit, contains, equals")
                elif kinds[0] == "per_unit":
                    if not _is_number(term["per_unit"]):
                        err(f"term {i}: 'per_unit' must be a number")
                elif not _is_number(term.get("add")):
                    err(f"term {i}: needs a numeric 'add'")

    elif spec_type == "categorical":
        need(spec.get("fact"), "'fact'")
        mapping = spec.get("map")
        if not (isinstance(mapping, dict) and mapping):
            err("'map' must be a non-empty object")
        else:
            for key, value in mapping.items():
                if _is_number(value):
                    unit(value, f"map['{key}']")
                elif isinstance(value, dict):
                    need(value.get("when"), f"map['{key}']")
                    unit(value.get("if_true"), f"map['{key}'].if_true")
                    unit(value.get("if_false"), f"map['{key}'].if_false")
                else:
                    err(f"map['{key}'] must be a number or a conditional object")

    elif spec_type == "boolean_map":
        need(spec.get("fact"), "'fact'")
        unit(spec.get("if_true"), "'if_true'")
        unit(spec.get("if_false"), "'if_false'")
        if "gate" in spec and not isinstance(spec["gate"], bool):
            err("'gate' must be true or false")


def attach_rule_specs(parameters: list[dict], specs: Any, errors: list[str]) -> None:
    if not isinstance(specs, dict):
        raise ConfigError("Rule specs file must be a JSON object keyed by parameter ID (p01, p02, ...).")
    known_ids = {p["id"] for p in parameters}
    for pid in sorted(set(specs) - known_ids):
        errors.append(f"Rule specs contain {pid}, which does not exist in the Excel document")
    for p in parameters:
        spec = specs.get(p["id"])
        if not isinstance(spec, dict):
            errors.append(f"{p['id']} ('{p['parameter']}'): no rule spec defined")
            continue
        if spec.get("parameter") != p["parameter"]:
            errors.append(
                f"{p['id']}: rule spec belongs to '{spec.get('parameter')}' but the Excel row is "
                f"'{p['parameter']}' (rows were probably reordered or renamed)"
            )
        body = {key: value for key, value in spec.items() if key != "parameter"}
        _validate_spec(p["id"], body, errors)
        p["rule_spec"] = body


# --------------------------------------------------------------------------
# 3. Intervals
# --------------------------------------------------------------------------
def attach_intervals(parameters: list[dict], intervals: Any, errors: list[str]) -> None:
    if not isinstance(intervals, dict):
        raise ConfigError("Intervals file must be a JSON object keyed by parameter ID (p01, p02, ...).")
    known_ids = {p["id"] for p in parameters}
    for pid in sorted(set(intervals) - known_ids):
        errors.append(f"Intervals contain {pid}, which does not exist in the Excel document")
    for p in parameters:
        p["interval"] = {"lower": None, "upper": None}
        entry = intervals.get(p["id"])
        if entry is None:
            continue
        if not isinstance(entry, dict):
            errors.append(f"{p['id']}: interval entry must be an object")
            continue
        if entry.get("parameter") != p["parameter"]:
            errors.append(
                f"{p['id']}: interval belongs to '{entry.get('parameter')}' but the Excel row is "
                f"'{p['parameter']}' (rows were probably reordered or renamed)"
            )
            continue
        lower, upper = entry.get("lower"), entry.get("upper")
        if lower is None and upper is None:
            continue  # not configured yet
        if lower is None or upper is None:
            errors.append(f"{p['id']}: provide both lower and upper, or neither")
        elif not (_is_number(lower) and _is_number(upper)):
            errors.append(f"{p['id']}: lower and upper must be numbers")
        elif not (0 <= lower <= 1 and 0 <= upper <= 1):
            errors.append(f"{p['id']}: lower and upper must be between 0 and 1 (got {lower}, {upper})")
        elif lower > upper:
            errors.append(f"{p['id']}: lower ({lower}) cannot be greater than upper ({upper})")
        else:
            p["interval"] = {"lower": lower, "upper": upper}


def missing_intervals(config: dict) -> list[str]:
    return [
        p["id"] for p in config["parameters"]
        if p["interval"]["lower"] is None or p["interval"]["upper"] is None
    ]


def require_complete(config: dict) -> None:
    """The scoring engine should call this before using a config."""
    missing = missing_intervals(config)
    if missing:
        raise ConfigError("Intervals are not configured for: " + ", ".join(missing))


def init_intervals(excel_path: Path | str = DEFAULT_EXCEL, intervals_path: Path | str = DEFAULT_INTERVALS) -> Path:
    """Create an empty intervals template. Never overwrites an existing file."""
    intervals_path = Path(intervals_path)
    if intervals_path.exists():
        raise ConfigError(f"{intervals_path} already exists; refusing to overwrite it.")
    parsed = parse_excel(excel_path)
    template = {p["id"]: {"parameter": p["parameter"], "lower": None, "upper": None} for p in parsed["parameters"]}
    intervals_path.parent.mkdir(parents=True, exist_ok=True)
    intervals_path.write_text(json.dumps(template, indent=2, ensure_ascii=False), encoding="utf-8")
    return intervals_path


# --------------------------------------------------------------------------
# 4. Build, save, load
# --------------------------------------------------------------------------
def build_config(
    excel_path: Path | str = DEFAULT_EXCEL,
    rule_specs_path: Path | str = DEFAULT_RULE_SPECS,
    intervals_path: Path | str = DEFAULT_INTERVALS,
    sheet_name: Optional[str] = None,
) -> dict:
    parsed = parse_excel(excel_path, sheet_name)
    parameters = parsed["parameters"]

    errors: list[str] = []
    attach_rule_specs(parameters, _read_json(rule_specs_path, "Rule specs file"), errors)
    attach_intervals(parameters, _read_json(intervals_path, "Intervals file"), errors)
    _raise_if(errors, "Configuration is invalid")

    config = {
        "config_version": _version_of(parameters),
        "generated_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "source": {
            "file": parsed["source_file"],
            "sha256": parsed["source_sha256"],
            "sheet": parsed["sheet"],
            "note": parsed["source_note"],
        },
        "total_weight": sum(p["weight"] for p in parameters),
        "parameters": parameters,
    }
    config["complete"] = not missing_intervals(config)
    return config


def save_config(
    config: dict,
    output_path: Path | str = DEFAULT_OUTPUT,
    versions_dir: Optional[Path | str] = DEFAULT_VERSIONS_DIR,
) -> Path:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    text = json.dumps(config, indent=2, ensure_ascii=False)
    output_path.write_text(text, encoding="utf-8")
    if versions_dir is not None:
        versions_dir = Path(versions_dir)
        versions_dir.mkdir(parents=True, exist_ok=True)
        archive = versions_dir / f"parameters_config_{config['config_version']}.json"
        if not archive.exists():
            archive.write_text(text, encoding="utf-8")
    return output_path


def load_config(path: Path | str = DEFAULT_OUTPUT, require_complete_intervals: bool = False) -> dict:
    """Load a generated config and verify it was not edited by hand afterwards."""
    config = _read_json(path, "Config file")
    parameters = config.get("parameters") if isinstance(config, dict) else None
    if not isinstance(parameters, list) or not parameters:
        raise ConfigError(f"Config file has no parameters: {path}")
    for p in parameters:
        for key in ("id", "parameter", "weight", "rule_spec", "interval"):
            if key not in p:
                raise ConfigError(f"Config file is missing '{key}' for a parameter: {path}")
    if _version_of(parameters) != config.get("config_version"):
        raise ConfigError(
            f"Config file {path} does not match its config_version; it was edited after generation. "
            "Re-run parser_config.py instead of editing it by hand."
        )
    if require_complete_intervals:
        require_complete(config)
    return config


# --------------------------------------------------------------------------
# Command line
# --------------------------------------------------------------------------
def main(argv: Optional[list[str]] = None) -> int:
    parser = argparse.ArgumentParser(description="Build parameters_config.json from the Excel rulebook.")
    parser.add_argument("--excel", default=str(DEFAULT_EXCEL))
    parser.add_argument("--rule-specs", default=str(DEFAULT_RULE_SPECS))
    parser.add_argument("--intervals", default=str(DEFAULT_INTERVALS))
    parser.add_argument("--output", default=str(DEFAULT_OUTPUT))
    parser.add_argument("--sheet", default=None, help="Sheet name (default: first sheet)")
    parser.add_argument("--strict", action="store_true", help="Fail if any interval is not configured")
    parser.add_argument("--init-intervals", action="store_true", help="Create an empty intervals template and exit")
    args = parser.parse_args(argv)

    try:
        if args.init_intervals:
            created = init_intervals(args.excel, args.intervals)
            print(f"Created empty interval template: {created}")
            return 0
        config = build_config(args.excel, args.rule_specs, args.intervals, args.sheet)
        if args.strict:
            require_complete(config)
        output = save_config(config, args.output)
    except ConfigError as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1

    by_field: dict[str, float] = {}
    for p in config["parameters"]:
        by_field[p["field"]] = by_field.get(p["field"], 0) + p["weight"]

    print(f"Config version : {config['config_version']}")
    print(f"Parameters     : {len(config['parameters'])}")
    print(f"Total weight   : {config['total_weight']}")
    for field, weight in by_field.items():
        print(f"  {field}: {weight}")
    print(f"Intervals set  : {'all' if config['complete'] else 'INCOMPLETE - missing ' + ', '.join(missing_intervals(config))}")
    print(f"Written to     : {output}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
