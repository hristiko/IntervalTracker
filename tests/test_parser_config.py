import pytest
from openpyxl import load_workbook

import parser_config as pc
from conftest import read_json, write_json


def build(workdir):
    return pc.build_config(workdir / "rulebook.xlsx", workdir / "rule_specs.json", workdir / "intervals.json")


def test_parses_all_11_parameters_with_stable_ids(workdir):
    parsed = pc.parse_excel(workdir / "rulebook.xlsx")
    assert [p["id"] for p in parsed["parameters"]] == [f"p{i:02d}" for i in range(1, 12)]


def test_source_text_is_preserved_including_typos(workdir):
    names = [p["parameter"] for p in pc.parse_excel(workdir / "rulebook.xlsx")["parameters"]]
    assert names[6] == "Revenu Trend"
    assert names[7] == "EBITDA Trenda"
    assert names[0] == "Date of incorporation"


def test_weights_and_total(workdir):
    params = pc.parse_excel(workdir / "rulebook.xlsx")["parameters"]
    assert [p["weight"] for p in params] == [6, 8, 5, 9, 7, 1, 6, 7, 8, 8, 7]
    assert sum(p["weight"] for p in params) == 72


def test_source_note_is_captured(workdir):
    note = pc.parse_excel(workdir / "rulebook.xlsx")["source_note"]
    assert note.startswith("Weigh and Rule can be customized")


def test_regression_no_crash_on_read_only_dimensions(workdir):
    assert len(pc.parse_excel(workdir / "rulebook.xlsx")["parameters"]) == 11


def test_build_is_complete_and_versioned(workdir):
    config = build(workdir)
    assert config["complete"] is True
    assert len(config["config_version"]) == 12
    assert config["total_weight"] == 72
    assert all("rule_spec" in p for p in config["parameters"])


def test_version_is_deterministic_and_changes_with_content(workdir):
    first, second = build(workdir), build(workdir)
    assert first["config_version"] == second["config_version"]
    intervals = read_json(workdir / "intervals.json")
    intervals["p09"]["lower"] = 0.60
    write_json(workdir / "intervals.json", intervals)
    assert build(workdir)["config_version"] != first["config_version"]


def test_save_and_load_roundtrip_and_archive(workdir):
    config = build(workdir)
    out = pc.save_config(config, workdir / "out.json", versions_dir=workdir / "versions")
    assert (workdir / "versions" / f"parameters_config_{config['config_version']}.json").exists()
    assert pc.load_config(out)["config_version"] == config["config_version"]


def test_hand_edited_config_is_detected(workdir):
    out = pc.save_config(build(workdir), workdir / "out.json", versions_dir=None)
    data = read_json(out)
    data["parameters"][0]["weight"] = 99
    write_json(out, data)
    with pytest.raises(pc.ConfigError, match="edited after generation"):
        pc.load_config(out)


@pytest.mark.parametrize("change, message", [
    ({"lower": 0.9, "upper": 0.5}, "cannot be greater"),
    ({"lower": 0.5, "upper": None}, "both lower and upper"),
    ({"lower": -0.1, "upper": 1}, "between 0 and 1"),
    ({"lower": 0.5, "upper": 1.5}, "between 0 and 1"),
    ({"lower": "low", "upper": 1}, "must be numbers"),
])
def test_invalid_intervals_are_rejected(workdir, change, message):
    intervals = read_json(workdir / "intervals.json")
    intervals["p03"].update(change)
    write_json(workdir / "intervals.json", intervals)
    with pytest.raises(pc.ConfigError, match=message):
        build(workdir)


def test_interval_for_unknown_id_is_rejected(workdir):
    intervals = read_json(workdir / "intervals.json")
    intervals["p99"] = {"parameter": "Ghost", "lower": 0, "upper": 1}
    write_json(workdir / "intervals.json", intervals)
    with pytest.raises(pc.ConfigError, match="p99"):
        build(workdir)


def test_reordered_excel_rows_are_caught_by_name_check(workdir):
    """Regression: IDs follow row order, so swapping rows must not silently shift intervals."""
    path = workdir / "rulebook.xlsx"
    workbook = load_workbook(path)
    sheet = workbook.worksheets[0]
    for column in range(1, 6):
        a, b = sheet.cell(3, column), sheet.cell(4, column)
        a.value, b.value = b.value, a.value
    workbook.save(path)
    with pytest.raises(pc.ConfigError, match="reordered or renamed"):
        build(workdir)


def test_unset_intervals_are_allowed_but_not_complete(workdir):
    intervals = read_json(workdir / "intervals.json")
    intervals["p05"].update({"lower": None, "upper": None})
    write_json(workdir / "intervals.json", intervals)
    config = build(workdir)
    assert config["complete"] is False
    assert pc.missing_intervals(config) == ["p05"]
    with pytest.raises(pc.ConfigError, match="p05"):
        pc.require_complete(config)


def test_missing_rule_spec_is_rejected(workdir):
    specs = read_json(workdir / "rule_specs.json")
    del specs["p11"]
    write_json(workdir / "rule_specs.json", specs)
    with pytest.raises(pc.ConfigError, match="p11.*no rule spec"):
        build(workdir)


def test_spec_referring_to_undeclared_fact_is_rejected(workdir):
    specs = read_json(workdir / "rule_specs.json")
    specs["p03"]["terms"][0]["fact"] = "contract_count_typo"
    write_json(workdir / "rule_specs.json", specs)
    with pytest.raises(pc.ConfigError, match="contract_count_typo"):
        build(workdir)


def test_unknown_rule_type_is_rejected(workdir):
    specs = read_json(workdir / "rule_specs.json")
    specs["p09"]["type"] = "magic"
    write_json(workdir / "rule_specs.json", specs)
    with pytest.raises(pc.ConfigError, match="unknown rule type"):
        build(workdir)


def test_non_contiguous_bands_are_rejected(workdir):
    specs = read_json(workdir / "rule_specs.json")
    specs["p01"]["bands"][1][0] = 15
    write_json(workdir / "rule_specs.json", specs)
    with pytest.raises(pc.ConfigError, match="contiguous"):
        build(workdir)


def test_missing_weight_is_rejected(workdir):
    path = workdir / "rulebook.xlsx"
    workbook = load_workbook(path)
    workbook.worksheets[0].cell(5, 5).value = None
    workbook.save(path)
    with pytest.raises(pc.ConfigError, match="weight is missing"):
        build(workdir)


def test_init_intervals_never_overwrites(workdir):
    with pytest.raises(pc.ConfigError, match="refusing to overwrite"):
        pc.init_intervals(workdir / "rulebook.xlsx", workdir / "intervals.json")
    template = pc.init_intervals(workdir / "rulebook.xlsx", workdir / "new_intervals.json")
    data = read_json(template)
    assert len(data) == 11 and all(v["lower"] is None for v in data.values())


def test_cli_strict_fails_when_incomplete(workdir, capsys):
    intervals = read_json(workdir / "intervals.json")
    intervals["p01"].update({"lower": None, "upper": None})
    write_json(workdir / "intervals.json", intervals)
    code = pc.main([
        "--excel", str(workdir / "rulebook.xlsx"), "--rule-specs", str(workdir / "rule_specs.json"),
        "--intervals", str(workdir / "intervals.json"), "--output", str(workdir / "out.json"), "--strict",
    ])
    assert code == 1
    assert not (workdir / "out.json").exists()


def test_interval_name_mismatch_is_rejected(workdir):
    """Intervals must follow their parameter, not just the p-number."""
    intervals = read_json(workdir / "intervals.json")
    intervals["p02"]["parameter"] = "Some other parameter"
    write_json(workdir / "intervals.json", intervals)
    with pytest.raises(pc.ConfigError, match="interval belongs to 'Some other parameter'"):
        build(workdir)


def test_rule_spec_name_mismatch_is_rejected(workdir):
    specs = read_json(workdir / "rule_specs.json")
    specs["p02"]["parameter"] = "Some other parameter"
    write_json(workdir / "rule_specs.json", specs)
    with pytest.raises(pc.ConfigError, match="rule spec belongs to 'Some other parameter'"):
        build(workdir)
